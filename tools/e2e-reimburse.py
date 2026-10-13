#!/usr/bin/env python3
"""
离线端到端测试：按《报销规范手册》整理报销材料（v0.19.0，用户 2026-10-09）。全程不连淘宝、不联网，订单、店铺、发票、报销人全部虚构
（报销人「张三」、抬头「某大学」）。

    env -u TMPDIR python3 tools/e2e-reimburse.py

需要 python3、playwright（含 Chromium）、openpyxl、python-docx、Pillow（读回生成的报销清单、用途说明、截图）。
流程：导入虚构订单表 → 第 2 步低值品标签（待确认排最前、点标签切换、记住）→ 没填姓名时点「整理报销文件」先打开设置 → 填姓名学号 →
「自动处理发票」：超过 1000 元的单自动打开模拟订单详情页分段截图，确认清单里 3D 打印订单的消息末尾追加索要明细的一句 →
发票表的颜色标签、「手动添加发票 / 附件」补支付记录 → 3D 打印明细从旺旺下载的活 → 整理报销文件（用户 2026-10-13 的格式：两个文件夹、
「序号_开票日期_金额-商品-数量」、附件「序号_类型」、报销清单.xlsx、用途说明 docx、截图）→ 报销记录（填到账、差额组合）→
导入已整理的发票文件夹认出历史批次（新旧两种文件夹名）、下一批序号接着编 → 批次进备份
"""
import csv
import io
import json
import os
import re
import shutil
import sys
import tempfile
import time
import zipfile
from pathlib import Path
from urllib.parse import urlsplit

from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parents[1]
TOOLS = ROOT / 'tools'
TITLE, TAX, NAME = '某大学', '121000009999999996', '张三'
BLANK = '<!doctype html><meta charset="utf-8"><body>（模拟空白页）</body>'
# 订单：key → (订单号, 日期, 店铺, 实付, 商品, 规格, 数量, 单价)；发票：key → (金额, 发票明细)
O = {
    'O1': ('5190000000000000401', '2026-08-01', '某某虚构接插件', '10.00', '杜邦线 公对母 40P', '20cm', 1, '10.00'),
    'O2': ('5190000000000000402', '2026-08-02', '某某虚构热缩管', '30.00', '热缩管 套装', '黑色', 1, '30.00'),         # 部分退款 5.00
    'O3': ('5190000000000000403', '2026-08-03', '某某虚构碳纤维加工', '1200.00', '碳纤维板 CNC 加工 3mm', '定制', 1, '1200.00'),
    'O4': ('5190000000000000404', '2026-08-04', '某某虚构仪表', '452.00', '数字万用表', '标准款', 1, '452.00'),
    'O5': ('5190000000000000405', '2026-08-05', '某某虚构衡器', '432.20', '电子秤', '0.01g', 1, '432.20'),
    'O6': ('5190000000000000406', '2026-08-06', '某某虚构3D打印', '80.00', '3D打印 光固化 手板 定制', '白色树脂', 1, '80.00'),
    'O7': ('5190000000000000407', '2026-08-07', '某某虚构模型', '66.00', '遥控车 配件', '套装', 1, '66.00'),
    'O8': ('5190000000000000408', '2026-08-08', '某某虚构杂货', '300.00', '某某说不清的器件', '一个', 1, '300.00'),
    'O9a': ('5190000000000000409', '2026-08-09', '某某虚构合开店', '10.00', '杜邦线 母对母 40P', '20cm', 1, '10.00'),
    'O9b': ('5190000000000000410', '2026-08-09', '某某虚构合开店', '15.00', '杜邦线 公对公 40P', '30cm', 1, '15.00'),
    'O10': ('5190000000000000411', '2026-08-10', '某某虚构排针', '6.00', '排针 2.54mm', '40P', 1, '6.00'),           # 没有发票
    'O11': ('5190000000000000412', '2026-08-11', '某某虚构部件', '350.00', '某某说不清的部件', '一个', 1, '350.00'),
}
NO = {k: v[0] for k, v in O.items()}
INV = {'O1': (10.5, '*电子元件*杜邦线'), 'O2': (25.0, '*塑料制品*热缩管'), 'O3': (1200.0, '*碳纤维制品*碳纤维板加工'), 'O4': (452.0, '*电子测量仪器*数字万用表'),
       'O5': (432.2, '*衡器*电子秤'), 'O6': (80.0, '*加工服务*3D打印加工费'), 'O7': (66.0, '*玩具*遥控车配件'), 'O8': (300.0, '*其他*器件'),
       'O9': (25.0, '*电子元件*杜邦线'), 'O11': (350.0, '*其他*部件')}
INVOICE_HTML = '''<!doctype html><meta charset="utf-8"><body style="font-family:sans-serif;padding:40px">
<h2>电子发票（普通发票）</h2><p>发票号码：{inv}</p><p>开票日期：{y}年{m}月{d}日</p>
<p>购买方信息 名称：某大学 统一社会信用代码/纳税人识别号：121000009999999996</p><p>销售方信息 名称：某某虚构商店</p>
<p>项目名称 {item} 金额 ¥{a1} 税额 ¥{a2}</p><p>价税合计（小写）¥{amt}</p></body>'''
fails = []


def check(ok, msg, detail=''):
    print(('  通过  ' if ok else '  失败  ') + msg + ('' if ok or not detail else '\n          ' + str(detail)[:600]))
    if not ok:
        fails.append(msg)


def wait_until(page, fn, timeout):
    end = time.time() + timeout
    while time.time() < end:
        v = fn()
        if v:
            return v
        page.wait_for_timeout(300)
    return None


def make_pdfs(p):
    """虚构发票 PDF（浏览器现场打印）：{key: PDF 字节}；O9a / O9b 合开一张"""
    b = p.chromium.launch()
    pg = b.new_page()
    out = {}
    for k, (amt, item) in INV.items():
        d = O['O9a' if k == 'O9' else k][1]
        inv = '2699000000000000' + (NO['O9a'] if k == 'O9' else NO[k])[-4:]
        y, m, dd = d.split('-')
        pg.set_content(INVOICE_HTML.format(inv=inv, y=y, m=m, d=f'{int(dd) + 2:02d}', item=item, amt=f'{amt:.2f}', a1=f'{amt * 0.9:.2f}', a2=f'{amt * 0.1:.2f}'))
        out[k] = pg.pdf()
    pg.set_content(INVOICE_HTML.format(inv='26990000000000000999', y='2026', m='03', d='01', item='*电子元件*示例商品', amt='100.00', a1='90.00', a2='10.00'))
    out['hist'] = pg.pdf()
    pg.set_content(INVOICE_HTML.format(inv='26990000000000000998', y='2026', m='04', d='20', item='*电子测量仪器*示例低值品', amt='88.00', a1='79.20', a2='8.80'))
    out['hist2'] = pg.pdf()
    b.close()
    return out


def write_csv(path):
    with open(path, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['订单号', '订单提交时间', '订单状态', '店铺名称', '商品名称', '型号款式', '商品数量', '商品金额', '实付金额', '运费'])
        for no, d, shop, pay, t, sku, q, price in O.values():
            w.writerow([no, d + ' 10:00:00', '交易成功', shop, t, sku, q, price, pay, '0.00'])


def run(p, tmp):
    pdfs = make_pdfs(p)
    dl_dir = tmp / '下载'
    (tmp / 'profile' / 'Default').mkdir(parents=True)
    (tmp / 'profile' / 'Default' / 'Preferences').write_text(json.dumps(
        {'download': {'default_directory': str(dl_dir), 'prompt_for_download': False, 'directory_upgrade': True}}), encoding='utf-8')
    ctx = p.chromium.launch_persistent_context(str(tmp / 'profile'), channel='chromium', headless=True, accept_downloads=True,
                                               env=dict(os.environ, HOME=str(tmp / 'home')), args=[
        '--disable-extensions-except=' + str(ROOT), '--load-extension=' + str(ROOT), '--no-proxy-server', '--host-resolver-rules=MAP * ~NOTFOUND'])
    detail_html = (TOOLS / 'mock-detail.html').read_text(encoding='utf-8')
    errors, opened = [], []

    def handle(route):
        url = route.request.url
        if url.startswith('https://trade.taobao.com/trade/detail/'):
            opened.append(url)
            return route.fulfill(status=200, content_type='text/html; charset=utf-8', body=detail_html)
        if url.startswith('https://i.taobao.com/') or url.startswith('https://market.m.taobao.com/'):
            return route.fulfill(status=200, content_type='text/html; charset=utf-8', body=BLANK)       # 刷新开票记录、读旺旺：空白页，让这两段超时跳过
        route.abort()
    ctx.route(lambda u: urlsplit(u).scheme in ('http', 'https', 'ws', 'wss'), handle)
    ctx.on('page', lambda pg: pg.on('pageerror', lambda e: errors.append(f'{pg.url[:60]}: {e}')))
    sw = ctx.service_workers[0] if ctx.service_workers else ctx.wait_for_event('serviceworker')
    eid = urlsplit(sw.url).hostname
    print('扩展已加载，ID', eid)

    print('\n[1] 导入虚构订单表，全部判为实验室；O2 部分退款 5.00')
    csv_path = tmp / '虚构订单表.csv'
    write_csv(csv_path)
    app = ctx.new_page()
    app.on('pageerror', lambda e: errors.append(f'主页: {e}'))
    ctx.new_cdp_session(app).send('Browser.setDownloadBehavior', {'behavior': 'default'})
    app.goto(f'chrome-extension://{eid}/index.html')
    app.set_input_files('#file', str(csv_path))
    app.wait_for_selector('#main:not([hidden])')
    app.evaluate('''([nos, o2]) => { const S = JSON.parse(localStorage.getItem('orderTriage.app.v1'));
        for (const o of S.orders) if (nos.includes(o.no)) for (const l of o.lines) S.decisions[l.key] = 'lab';
        const o = S.orders.find(o => o.no === o2); S.refundAmt = { [o.lines[0].key]: 5 };
        localStorage.setItem('orderTriage.app.v1', JSON.stringify(S)); }''', [list(NO.values()), NO['O2']])
    app.reload(); app.wait_for_selector('#main:not([hidden])'); app.wait_for_timeout(500)
    S = lambda: app.evaluate("JSON.parse(localStorage.getItem('orderTriage.app.v1'))")
    store = lambda k: app.evaluate('k => chrome.storage.local.get(k).then(r => r[k])', k)

    print('\n[2] 第 2 步「核对商品」：低值品标签（单价超过 200 元）；判断不出的黄色「是否低值品？」排最前，点标签切换并记住')
    app.click('.flow li[data-step="1"]'); app.wait_for_timeout(300)
    app.click('#cat-seg [data-cat="all"]'); app.wait_for_timeout(200)       # 分类筛选选「全部」：实验室、个人一起看
    pills = app.evaluate('''() => Object.fromEntries([...document.querySelectorAll('.line')].map(l => [l.querySelector('.title').textContent.trim(),
        l.querySelector('.low-ask') ? '是否低值品？' : (l.querySelector('[data-low]') || {}).textContent || '']))''')
    check(pills.get('数字万用表') == '低值品' and pills.get('电子秤') == '低值品', '数字万用表 452、电子秤 432.20：深色「低值品」', pills)
    check(pills.get('碳纤维板 CNC 加工 3mm') == '非低值品' and pills.get('杜邦线 公对母 40P') == '', '碳纤维板 1200（耗材词）：「非低值品」；200 元以下的没有标签', pills)
    check(pills.get('某某说不清的器件') == '是否低值品？' and pills.get('某某说不清的部件') == '是否低值品？', '关键词判断不出的：黄色「是否低值品？」', pills)
    asked = app.evaluate("[...document.querySelectorAll('.line')].filter(l => l.querySelector('.title').textContent.includes('说不清')).map(l => l.className)")
    check(len(asked) == 2 and all('is-uns' in x for x in asked), '待确认是否低值品的整行标黄（和待定一样，位置按下单日期）', asked)
    check(app.is_disabled('#summary [data-flow="sort-done"]') and '是否低值品' in app.inner_text('#summary .fd-hint'),
          '还有待确认的时「确认核对完成」不能点，提示一行点黄色标签', app.inner_text('#summary .fd-hint'))
    asks = app.evaluate('''() => [...document.querySelector('.line .low-ask').querySelectorAll('button')].map(b => b.textContent + '|' + (b.title ? 'tip' : ''))''')
    check(asks == ['是|tip', '否|tip'], '待确认时直接给「是 / 否」两个选项，都有悬停说明', asks)
    app.click('.line:has-text("某某说不清的器件") [data-low][data-lowv="low"]'); app.wait_for_timeout(300)
    t1 = app.inner_text('.line:has-text("某某说不清的器件") [data-low]')
    app.click('.line:has-text("某某说不清的器件") [data-low]'); app.wait_for_timeout(300)
    t2 = app.inner_text('.line:has-text("某某说不清的器件") [data-low]')
    lv = S().get('lowval') or {}
    check(t1 == '低值品' and t2 == '非低值品' and list(lv.values()) == ['no'], '点「是」变「低值品」，再点标签变「非低值品」；按商品标题记住（S.lowval）', (t1, t2, lv))

    print('\n[3] 没填姓名时点「整理报销文件」：先打开设置，顶上一行提示，不选文件夹；填好姓名（没有学号）')
    app.click('.flow li[data-step="3"]'); app.wait_for_timeout(300)
    choosers = []
    app.on('filechooser', lambda fc: choosers.append(fc))
    app.click('#summary [data-flow="pack"]'); app.wait_for_timeout(500)
    check(app.is_visible('#dlg-settings') and '姓名' in app.inner_text('#set-note') and not choosers, '打开了设置，提示「请先填写报销人姓名」，没有弹出选文件夹', app.inner_text('#set-note'))
    check(app.locator('#person-sid').count() == 0 and '学号' not in app.inner_text('#dlg-settings'), '设置里没有学号')
    app.fill('#person-name', NAME); app.fill('#inv-title', TITLE); app.fill('#inv-tax', TAX)
    app.click('#rules-save'); app.wait_for_selector('#dlg-settings', state='hidden')
    check(S().get('person') == {'name': NAME}, '姓名存进本机设置', S().get('person'))
    app.click('#summary [data-flow="pack"]'); app.wait_for_timeout(500)
    check(len(choosers) == 1, '填好后再点：直接选择发票文件夹', len(choosers))

    print('\n[4] 下载好的发票（除 O6、O10）；「自动处理发票」：超过 1000 元的单自动截订单页面；3D 打印订单的消息末尾追加一句')
    folder = tmp / '订单分拣-发票'
    folder.mkdir()
    now = int(time.time() * 1000)
    dl = {}

    def put(k, data):
        no, d, shop, pay = O[k][:4]
        name = f'{d}_{float(pay):g}_{shop}_{no}.pdf'
        (folder / name).write_bytes(data)
        dl[no] = [{'file': name, 'path': str(folder / name), 'at': now, 'from': 'platform', 'src': ''}]
    for k in ('O1', 'O2', 'O3', 'O4', 'O5', 'O7', 'O8', 'O11'):
        put(k, pdfs[k])
    put('O9a', pdfs['O9']); put('O9b', pdfs['O9'])
    app.evaluate('f => chrome.storage.local.set({ dlDone: f })', dl)
    app.wait_for_timeout(800)
    app.click('.flow li[data-step="2"]'); app.wait_for_timeout(400)
    app.evaluate('() => { window.__otDev.tmo = { sync: 1500, scan: 1500 }; }')
    app.click('#summary [data-flow="inv-run"]')
    shown = wait_until(app, lambda: app.locator('#dlg-list[open]').count(), 120)
    msgs = app.evaluate("[...document.querySelectorAll('#list-rows .ask-msg')].map(x => x.textContent)") if shown else []
    exp = app.evaluate('''([o6, o10, t, x]) => { const I = window.Invoice, RB = window.Reimburse;
        const one = (no, d, a) => I.renderMsg(I.DEFAULT_TEMPLATE, { orders: [{ no, date: d, amount: a }], title: t, taxId: x, email: '' });
        return [one(o6, '2026-08-06', 80), RB.ASK_3D, one(o10, '2026-08-10', 6)]; }''', [NO['O6'], NO['O10'], TITLE, TAX])
    m6 = next((m for m in msgs if NO['O6'] in m), '')
    m10 = next((m for m in msgs if NO['O10'] in m), '')
    check(bool(shown) and m6 == exp[0] + exp[1], '确认清单里 3D 打印订单的消息 = 原消息（一字不改）+ 末尾「另外麻烦提供这单的 3D 打印明细清单…」', (m6, exp[0] + exp[1]))
    check(m10 == exp[2] and exp[1] not in m10, '其他订单的消息不变、不追加', m10)
    app.click('#list-cancel')
    fin = wait_until(app, lambda: '发票处理完成' in app.inner_text('#summary') and app.inner_text('#summary'), 60)
    check(bool(fin), '取消后「自动处理发票」结束')
    att3 = app.evaluate('no => __otDev.att(no)', NO['O3'])
    check(any(a['kind'] == '订单页面' for a in att3), '超过 1000 元的 O3：自动打开订单详情页截了「订单页面」（存在本机 IndexedDB）', att3)
    check(not [u for u in opened if NO['O1'] in u], '不需要截图的单（O1）没有被打开详情页', [u[-25:] for u in opened])
    pay = (S().get('payInfo') or {}).get(NO['O3'], {})
    check(pay.get('alipay') == '2026090122001400000000000403' and pay.get('paidAt') == '2026-09-01 10:00:05', '详情页上的支付宝交易号、付款时间记下来（写进 README.txt）', pay)
    check(not [t for t in app.context.pages if '/trade/detail/' in t.url], '截完图详情页都关掉了', [t.url for t in app.context.pages])

    print('\n[5] 发票表的颜色标签：低值品（深色）+「需先开低值票」、缺材料（橙色）；「手动添加发票 / 附件」补支付记录')
    app.click('.flow li[data-step="2"]'); app.wait_for_timeout(500)

    def row(k):
        return app.evaluate('''no => { const tr = [...document.querySelectorAll('.inv-table tbody tr')].find(tr => tr.textContent.includes(no)); if (!tr) return {};
            return { tags: [...tr.querySelectorAll('.tags .tag')].map(t => t.textContent + '|' + t.className), detail: tr.children[5].textContent,
                     add: (tr.querySelector('label.add') || {}).textContent || '' }; }''', NO[k])
    r4, r3 = row('O4'), row('O3')
    check(any(t.startswith('低值品|') and 't-low' in t for t in r4.get('tags', [])) and '需先开低值票' in r4.get('detail', ''), 'O4 数字万用表：深色「低值品」，下面一行「需先开低值票」', r4)
    check(any(t.startswith('需补支付记录|') and 't-mat' in t for t in r3.get('tags', [])) and not any('需订单截图' in t for t in r3.get('tags', []))
          and r3.get('add') == '手动添加发票 / 附件', 'O3（超过 1000 元，已截图）：橙色「需补支付记录」；次要链接变成「手动添加发票 / 附件」', r3)
    o11 = app.evaluate('''no => { const tr = [...document.querySelectorAll('.inv-table tbody tr')].find(tr => tr.textContent.includes(no)); const a = tr && tr.querySelector('.low-ask'); return a ? a.textContent : ''; }''', NO['O11'])
    check(o11 == '是否低值品？是否', '待确认的 O11：发票表里也是黄色「是否低值品？」加「是 / 否」', o11)
    app.set_input_files(f'input[data-inv="attach"][data-no="{NO["O3"]}"]', str(TOOLS / 'fixtures' / 'photo-mock.png'))
    got = wait_until(app, lambda: (a := app.evaluate('no => __otDev.att(no)', NO['O3'])) and any(x['kind'] == '支付记录' for x in a) and a, 10)
    check(bool(got), '选一张截图：自动认作这单缺的「支付记录」附件', got)
    app.wait_for_timeout(400)
    check(not any('支付记录' in t for t in row('O3').get('tags', [])), '补上后 O3 不再标「需补支付记录」', row('O3'))

    print('\n[6] 3D 打印订单的明细：卖家在旺旺发来的表格会排进下载、挂成「3D打印明细」附件（不算发票）')
    put('O6', pdfs['O6'])
    app.evaluate('f => chrome.storage.local.get("dlDone").then(r => chrome.storage.local.set({ dlDone: Object.assign({}, r.dlDone, f) }))', {NO['O6']: dl[NO['O6']]})
    app.evaluate('''([no, shop]) => chrome.storage.local.set({ chatScan: { at: Date.now(), convs: { [shop]: { at: Date.now(), orders: [no], first: '2026-08-06 10:00:00',
        asks: [{ time: '2026-08-07 10:00:00', text: '需要发票 ' + no, nos: [no] }], files: [], images: [], email: [], cards: [],
        docs: [{ time: '2026-08-08 10:00:00', name: '3D打印明细.xlsx', size: '9 KB' }] } } } })''', [NO['O6'], O['O6'][2]])
    app.wait_for_timeout(800)
    jobs = app.evaluate('nos => __otDev.jobsFor(nos)', [NO['O6']])
    check(jobs == [NO['O6'] + ' 3D打印明细.xlsx'], '缺 3D 打印明细的 O6：卖家发的 3D打印明细.xlsx 排进下载（带附件类型）', jobs)
    check(any(t.startswith('需3D打印明细') for t in row('O6').get('tags', [])), 'O6 还没拿到明细：橙色「需3D打印明细」', row('O6'))

    print('\n[7] 整理报销文件（用户 2026-10-13 的格式）：预览、待确认低值品先确认、生成「YYMMDD_第1批采购报销_金额」「…低值品_金额」两个文件夹')
    app.click('.flow li[data-step="3"]'); app.wait_for_timeout(300)
    app.set_input_files('#inv-pack-dir', str(folder))
    app.wait_for_selector('#dlg-pack[open]', timeout=90000)
    note = app.inner_text('#pack-note')
    today = time.strftime('%y%m%d')
    check(f'{today}_第1批采购报销_' in note and '待确认是否低值品' in note and app.is_disabled('#pack-go'), '预览写明文件夹名；有待确认低值品时「生成」不能点', note)
    att7 = app.evaluate('no => __otDev.att(no)', NO['O7'])
    check(any(a['kind'] == '订单页面' for a in att7), '发票明细含「玩具」的 O7：选文件夹时自动截了订单页面（要附在用途说明文末）', att7)
    app.click(f'#pack-list [data-lowno="{NO["O11"]}"]'); app.wait_for_timeout(500)
    check(not app.is_disabled('#pack-go') and app.input_value('#pack-name') == '第 1 批' and app.input_value('#pack-start') == '1',
          '点黄色标签确认 O11 为低值品后可以生成；批次名默认「第 1 批」，起始序号默认 1', (app.input_value('#pack-name'), app.input_value('#pack-start')))
    prev_js = """() => [...document.querySelectorAll('#pack-list .pk-row')].map(r => ({ seq: r.querySelector('.pk-seq').textContent, no: r.dataset.no,
        file: r.querySelector('.detail').textContent.split('　←')[0], dir: r.querySelector('.detail').title.split('/')[0], main: r.querySelector('.pk-main').innerText }))"""
    app.fill('#pack-start', '101'); app.wait_for_timeout(300)
    p101 = app.evaluate(prev_js)
    check(p101 and p101[0]['seq'] == '101' and p101[0]['file'].startswith('101_'), '起始序号改成 101：预览从 101 编起', p101[:1])
    app.fill('#pack-start', '1'); app.wait_for_timeout(300)
    prev = app.evaluate(prev_js)
    by = {x['no']: x for x in prev}
    want = {'O1': '1_260803_10.50-杜邦线公对母40P-1件.pdf', 'O2': '2_260804_25.00-热缩管套装-1件.pdf', 'O3': '3_260805_1200.00-碳纤维板CNC加工3mm-1件.pdf',
            'O4': '4_260806_452.00-数字万用表-1件.pdf', 'O5': '5_260807_432.20-电子秤-1件.pdf', 'O7': '7_260809_66.00-遥控车配件-1件.pdf',
            'O9a': '9+10_260811_25.00-杜邦线母对母40P等-2单合开.pdf', 'O11': '11_260813_350.00-某某说不清的部件-1件.pdf'}
    bad = {k: by.get(NO[k], {}).get('file') for k, v in want.items() if by.get(NO[k], {}).get('file') != v}
    check(not bad, '发票文件名「序号_开票日期_金额-商品摘要-数量」：按下单顺序编号（低值品、普通共用一套），合开的「9+10_…等-2单合开」', bad)
    lowd, normd = by.get(NO['O4'], {}).get('dir', ''), by.get(NO['O1'], {}).get('dir', '')
    check(normd == f'{today}_第1批采购报销_1706.50' and lowd == f'{today}_第1批采购报销低值品_1234.20' and all(by.get(NO[k], {}).get('dir') == lowd for k in ('O4', 'O5', 'O11')),
          '普通发票、低值品各一个文件夹：「YYMMDD_第1批采购报销_1706.50」「YYMMDD_第1批采购报销低值品_1234.20」', (normd, lowd))
    check('缺：3D打印明细' in by.get(NO['O6'], {}).get('main', '') and '缺：用途说明' in by.get(NO['O7'], {}).get('main', ''), '缺材料的标出「缺：3D打印明细」「缺：用途说明」', (by.get(NO['O6']), by.get(NO['O7'])))
    app.click('#pack-go')
    out = dl_dir / '订单分拣-报销'
    zps = wait_until(app, lambda: out.is_dir() and len(z := sorted(x for x in out.iterdir() if x.suffix == '.zip')) == 2 and z, 60) or []
    app.wait_for_timeout(1500)
    total = 10.5 + 25 + 1200 + 452 + 432.2 + 80 + 66 + 300 + 25 + 350
    check([z.name for z in zps] == sorted([normd + '.zip', lowd + '.zip']), '两个文件夹各一个同名压缩包', [z.name for z in zps])
    zs = {z.stem: zipfile.ZipFile(z) for z in zps}
    names = {k: sorted(n[len(k) + 1:] for n in z.namelist()) for k, z in zs.items()}
    nn, ln = names.get(normd, []), names.get(lowd, [])
    check(nn == sorted(['1_260803_10.50-杜邦线公对母40P-1件.pdf', '2_260804_25.00-热缩管套装-1件.pdf', '3_260805_1200.00-碳纤维板CNC加工3mm-1件.pdf', '3_支付记录.png', '3_订单页面.jpg',
                        '6_260808_80.00-3D打印光固化手板定制-1件.pdf', '7_260809_66.00-遥控车配件-1件.pdf', '7_用途说明.docx', '7_订单页面.jpg',
                        '8_260810_300.00-某某说不清的器件-1件.pdf', '9+10_260811_25.00-杜邦线母对母40P等-2单合开.pdf', '报销清单.xlsx']),
          '普通发票文件夹：发票、附件（「序号_订单页面.jpg」「序号_支付记录.png」「序号_用途说明.docx」）和报销清单.xlsx 放在一起，没有子文件夹、README.txt', nn)
    check(ln == sorted(['4_260806_452.00-数字万用表-1件.pdf', '5_260807_432.20-电子秤-1件.pdf', '11_260813_350.00-某某说不清的部件-1件.pdf', '报销清单.xlsx']),
          '低值品文件夹：3 张低值品和自己的报销清单.xlsx', ln)
    check(all((out / k / n).is_file() for k, ns in names.items() for n in ns), '下载文件夹里的文件夹和压缩包内容一样',
          [k + '/' + n for k, ns in names.items() for n in ns if not (out / k / n).is_file()][:3])
    check(not [d for d in out.iterdir() if re.search('报销给|' + NAME + '|科创|海创', d.name)], '文件夹、压缩包名里没有报销给谁、姓名、单位名', [d.name for d in out.iterdir()])
    z = zs.get(normd)
    if z:
        from PIL import Image
        im = Image.open(io.BytesIO(z.read(normd + '/3_订单页面.jpg')))
        check(im.height > 2000, f'订单页面截图是分段截取后拼接的长图（{im.size[0]}×{im.size[1]}，模拟页约 3100 像素高）', im.size)
        rgb = im.convert('RGB')
        ys = [y for y in range(0, im.height, 4) if (lambda c: c[0] > 200 and c[1] < 150 and c[2] < 90)(rgb.getpixel((1000, y)))]
        check(ys and min(ys) < 10 and max(ys) < 60, '固定导航栏（橙色）只在第一屏顶上出现一次，后面几段截图前藏起来了', ys[:5] + ys[-5:])
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(z.read(normd + '/报销清单.xlsx')))
        rows = [[c.value for c in r] for r in wb['报销清单'].iter_rows()]
        check(rows[0] == ['序号', '文件名', '开票日期', '金额', '发票号码', '订单号', '商品'], '报销清单.xlsx（openpyxl 读回）：序号、文件名、开票日期、金额、发票号码、订单号、商品', rows[0])
        r3 = next((r for r in rows if r[0] == 3), [])
        r9 = next((r for r in rows if r[0] == '9+10'), [])
        check(r3 and r3[1] == '3_260805_1200.00-碳纤维板CNC加工3mm-1件.pdf' and r3[2] == '2026-08-05' and r3[3] == 1200 and r3[4] == '26990000000000000403' and r3[5] == NO['O3'],
              '超过 1000 元那行：序号、文件名、开票日期、金额（数值）、发票号码、订单号', r3)
        check(r9 and NO['O9a'] in r9[5] and NO['O9b'] in r9[5] and r9[3] == 25, '合开那行：序号「9+10」、两个订单号', r9)
        check(rows[-1][0] == '合计' and abs(rows[-1][3] - 1706.5) < 0.005 and len(rows) == 9, '普通发票 7 张 + 表头 + 合计行', rows[-1])
        check('尚无发票' in wb.sheetnames and any(NO['O10'] in str(c.value) for r in wb['尚无发票'].iter_rows() for c in r), '第二张表「尚无发票」列出 O10', wb.sheetnames)
        lw = openpyxl.load_workbook(io.BytesIO(zs[lowd].read(lowd + '/报销清单.xlsx')))
        lr = [[c.value for c in r] for r in lw['报销清单'].iter_rows()]
        check([r[0] for r in lr[1:-1]] == [4, 5, 11] and abs(lr[-1][3] - 1234.2) < 0.005, '低值品的报销清单：序号 4、5、11，合计 1234.20', lr)
        import docx
        d = docx.Document(io.BytesIO(z.read(normd + '/7_用途说明.docx')))
        txt = '\n'.join(x.text for x in d.paragraphs)
        check(d.paragraphs[0].text == '用途说明' and f'订单号：{NO["O7"]}' in txt and '下单日期：2026-08-07' in txt and '用途：____' in txt and '金额：66.00 元' in txt
              and f'报销人：{NAME}' in txt and len(d.inline_shapes) >= 1, '用途说明 docx（python-docx 读回）：标题、订单号、下单日期、金额、「用途：____」留空、报销人，文末内嵌订单截图', (txt[:200], len(d.inline_shapes)))
    for z in zs.values():
        z.close()

    print('\n[8] 报销记录：整理后自动记「第 1 批」；订单显示「已整理（第 1 批）」；填到账（分期）→ 差额组合 → 已到账')
    app.wait_for_timeout(500)
    bs = app.evaluate('__otDev.batches()')
    b = bs[0] if bs else {}
    check(len(bs) == 1 and b.get('name') == '第 1 批' and b.get('n') == 10 and abs(b.get('amount', 0) - total) < 0.005 and abs(b.get('low', 0) - 1234.2) < 0.005
          and len(b.get('invNos', [])) == 10 and len(b.get('nos', [])) == 11 and b.get('label') == '待到账' and b.get('last') == 11,
          '批次：名称、张数、金额、低值品金额、发票号、订单号、最后一个序号 11；状态「待到账」', b)
    check(app.evaluate('__otDev.packStart()') == 12, '下一批的起始序号默认 12（接着这一批）', app.evaluate('__otDev.packStart()'))
    app.click('.flow li[data-step="2"]'); app.wait_for_timeout(400)
    app.evaluate("document.querySelectorAll('details.inv-sect').forEach(d => d.open = true)"); app.wait_for_timeout(200)
    st1 = app.evaluate("no => { const tr = [...document.querySelectorAll('.inv-table tbody tr')].find(tr => tr.innerText.includes(no)); return tr ? tr.querySelector('.st').textContent : ''; }", NO['O1'])
    check(st1 == '已整理（第 1 批）', '这批的订单显示「已整理（第 1 批）」', st1)
    app.click('.flow li[data-step="3"]'); app.wait_for_timeout(300)
    br = app.inner_text('#batches .b-row')
    check('第 1 批' in br and '待到账' in br and app.locator('#summary .flow-acts [data-pay]').count() == 0, '第 4 步说明区下方「报销记录」一行：第 1 批 · 待到账；「填写到账」不在主按钮区', br)
    check('tone-wait' in app.get_attribute('#batches .b-row .tag', 'class'), '待到账是黄色标签')
    app.click('#batches [data-pay]'); app.wait_for_selector('#dlg-pay[open]')
    app.fill('#pay-amt', f'{total - 884.2:.2f}'); app.fill('#pay-date', '2026-10-01'); app.click('#pay-ok'); app.wait_for_timeout(400)
    hint = app.inner_text('#batches .b-hint') if app.locator('#batches .b-hint').count() else ''
    check('有差额' in app.inner_text('#batches .b-row') and 'tone-bad' in app.get_attribute('#batches .b-row .tag', 'class')
          and hint == '差额可能是：4 号 数字万用表 452.00 + 5 号 电子秤 432.20', '到账少 884.20：红色「有差额」，一行写出差额可能是哪两张低值品', hint)
    app.click('#batches [data-pay]'); app.wait_for_selector('#dlg-pay[open]')
    check(app.input_value('#pay-amt') == '884.20', '再填时默认是剩下的差额', app.input_value('#pay-amt'))
    app.fill('#pay-date', '2026-10-08'); app.click('#pay-ok'); app.wait_for_timeout(400)
    b = app.evaluate('__otDev.batches()')[0]
    check(b.get('label') == '已到账' and len(b.get('pays', [])) == 2 and 'tone-ok' in app.get_attribute('#batches .b-row .tag', 'class'), '分两次到账：绿色「已到账」', b.get('pays'))

    print('\n[9] 「更多 → 导入已整理的发票文件夹」：「YYMMDD_……第X批……_金额_报销给某人」「YYMMDD_第N批采购报销低值品_金额」都认成历史批次；序号接着编')
    hist = tmp / '已整理' / '260401_实验室第一批采购_100.00_报销给张三'
    hist.mkdir(parents=True)
    (hist / '1_260301_100.00-示例商品-1件.pdf').write_bytes(pdfs['hist'])
    (tmp / '已整理' / '说明.pdf').write_bytes(pdfs['hist'])
    hist2 = tmp / '已整理' / '260501_第3批采购报销低值品_88.00'
    hist2.mkdir(parents=True)
    (hist2 / '205_260420_88.00-示例低值品-1件.pdf').write_bytes(pdfs['hist2'])
    app.set_input_files('#inv-have-dir', str(tmp / '已整理'))
    imp = lambda: [b for b in app.evaluate('__otDev.batches()') if b.get('src') == 'import']
    got = wait_until(app, lambda: (x := imp()) and len(x) == 2 and x, 30) or imp()
    hb = next((b for b in got if b.get('name') == '第一批'), {})
    h3 = next((b for b in got if b.get('name') == '第3批'), {})
    check(hb.get('date') == '2026-04-01' and hb.get('amount') == 100 and hb.get('n') == 1 and hb.get('label') == '待填到账',
          '认出历史批次：名称「第一批」、日期 2026-04-01、金额 100、1 张，状态「待填到账」', hb)
    check(h3.get('date') == '2026-05-01' and h3.get('amount') == 88 and abs(h3.get('low', 0) - 88) < 0.005 and h3.get('dir') == '260501_第3批采购报销低值品_88.00',
          '插件格式的低值品文件夹「260501_第3批采购报销低值品_88.00」也认出来（第3批、低值品 88.00）', h3)
    check(app.evaluate('__otDev.packStart()') == 206, '导入的发票文件名里最大序号 205：下一批从 206 编起', app.evaluate('__otDev.packStart()'))
    app.wait_for_timeout(500)
    check('待填到账' in app.inner_text('#batches'), '报销记录里多了这一批（待填到账）', app.inner_text('#batches'))

    print('\n[10] 批次进备份')
    app.click('#more summary'); app.click('#data-backup')
    bk = wait_until(app, lambda: (d := dl_dir / '订单分拣-备份').is_dir() and (f := [x for x in d.iterdir() if x.suffix == '.json']) and f[0], 15)
    data = json.loads(bk.read_text(encoding='utf-8')) if bk else {}
    st = json.loads((data.get('localStorage') or {}).get('orderTriage.app.v1', '{}'))
    check(len(st.get('batches', [])) == 3 and len(st['batches'][0].get('pays', [])) == 2, '备份文件里有三个批次和到账记录', [b.get('name') for b in st.get('batches', [])])
    check(not errors, '页面没有报错', errors[:3])
    ctx.close()


def main():
    tmp = Path(tempfile.mkdtemp(prefix='order-triage-rb-'))
    try:
        with sync_playwright() as p:
            run(p, tmp)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)        # 连同测试浏览器的配置目录（localStorage、扩展存储、IndexedDB、下载的假发票）一起删掉
    print('\n' + ('全部通过' if not fails else f'{len(fails)} 项失败：\n  ' + '\n  '.join(fails)))
    sys.exit(1 if fails else 0)


if __name__ == '__main__':
    main()
